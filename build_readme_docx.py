"""将 README.md 转为 README.docx（简易 Markdown 子集）。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "README.md"
DOCX_PATH = ROOT / "README.docx"


def _set_run_font(run, *, size_pt: float = 11, bold: bool = False, code: bool = False) -> None:
    run.font.size = Pt(size_pt)
    run.bold = bold
    name = "Consolas" if code else "Times New Roman"
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    r_fonts = rpr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), "宋体")
    if code:
        run.font.color.rgb = RGBColor(0x1C, 0x4D, 0x6E)


def _add_inline_runs(paragraph, text: str, *, size_pt: float = 11) -> None:
    """支持 **粗体**、`代码`、[文字](url) 的简单内联。"""
    pattern = re.compile(
        r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    )
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            _set_run_font(run, size_pt=size_pt)
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, size_pt=size_pt, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, size_pt=size_pt, code=True)
        else:
            link_m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            assert link_m
            label, url = link_m.group(1), link_m.group(2)
            run = paragraph.add_run(f"{label} ({url})")
            _set_run_font(run, size_pt=size_pt)
            run.font.color.rgb = RGBColor(0x0F, 0x6A, 0x5C)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_font(run, size_pt=size_pt)


def _add_heading(doc: Document, text: str, level: int) -> None:
    text = text.strip()
    p = doc.add_paragraph()
    sizes = {1: 18, 2: 14, 3: 12}
    run = p.add_run(text)
    _set_run_font(run, size_pt=sizes.get(level, 12), bold=True)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12 if level > 1 else 0)
    p.paragraph_format.space_after = Pt(6)


def _add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _add_inline_runs(p, text.strip())
    p.paragraph_format.space_after = Pt(4)


def _add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        _set_run_font(run, size_pt=9, code=True)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            # 单元格内去掉 markdown 粗体标记的简单处理
            plain = re.sub(r"\*\*([^*]+)\*\*", r"\1", cell_text.strip())
            plain = re.sub(r"`([^`]+)`", r"\1", plain)
            run = p.add_run(plain)
            _set_run_font(run, size_pt=10, bold=(i == 0))
    doc.add_paragraph()


def _parse_table_block(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if re.match(r"^\|?\s*:?-+:?\s*\|", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells:
            rows.append(cells)
    return rows


def md_to_docx(md_text: str, out_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("```"):
            fence = stripped[3:].strip()
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # closing ```
            _add_code_block(doc, block)
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(doc, _parse_table_block(block))
            continue

        if stripped.startswith("# "):
            _add_heading(doc, stripped[2:], 1)
            i += 1
            continue
        if stripped.startswith("## "):
            _add_heading(doc, stripped[3:], 2)
            i += 1
            continue
        if stripped.startswith("### "):
            _add_heading(doc, stripped[4:], 3)
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            _add_inline_runs(p, stripped[2:])
            p.paragraph_format.left_indent = Cm(0.5)
            i += 1
            continue

        if re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, stripped[2:])
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            text = re.sub(r"^\d+\.\s+", "", stripped)
            _add_inline_runs(p, text)
            i += 1
            continue

        _add_para(doc, stripped)
        i += 1

    doc.save(out_path)


def main() -> None:
    md = MD_PATH.read_text(encoding="utf-8")
    md_to_docx(md, DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()
