"""生成未掌握单词的中英释义 Word 文档。"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _set_run_font(run, *, size_pt: float = 12) -> None:
    """小四号 = 12 磅。"""
    run.font.size = Pt(size_pt)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    r_fonts = rpr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")


def build_vocabulary_docx(
    pairs: list[tuple[str, str]],
    *,
    title: str = "未掌握单词释义表",
    english_header: str = "英文",
    chinese_header: str = "释义",
    empty_message: str = "本次没有需要导出的未掌握单词。",
) -> bytes:
    """
    生成两列表格文档：左列英文，右列释义（可含多行）。
    每行一对，字号小四（12pt）。

    pairs 的右列可以是 ECDICT 多行释义，也可以是百度翻译短句；
    此处可以改造接入百度翻译 API：由调用方传入 translate_words 的结果即可。
    """
    document = Document()
    section = document.sections[0]
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    _set_run_font(run, size_pt=14)
    run.bold = True

    if not pairs:
        tip = document.add_paragraph()
        run = tip.add_run(empty_message)
        _set_run_font(run)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    for cell, text in zip(header, (english_header, chinese_header), strict=True):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(text)
        _set_run_font(run)
        run.bold = True

    for english, definition in pairs:
        row = table.add_row().cells
        # 英文
        row[0].text = ""
        p0 = row[0].paragraphs[0]
        run = p0.add_run(english)
        _set_run_font(run)
        # 释义：按换行拆成多个段落，便于阅读
        row[1].text = ""
        lines = (definition or "").splitlines() or [""]
        first = True
        for line in lines:
            if first:
                paragraph = row[1].paragraphs[0]
                first = False
            else:
                paragraph = row[1].add_paragraph()
            run = paragraph.add_run(line)
            _set_run_font(run)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
