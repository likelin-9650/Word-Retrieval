"""根据词库为大词典分类结果着色。"""

from __future__ import annotations

import io
import re
from typing import Any, Iterable

from docx import Document
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

from word_retrieval.extractor import iter_english_tokens

GREEN = RGBColor(0x1B, 0x8F, 0x3A)
YELLOW = RGBColor(0xC4, 0xA0, 0x00)
RED = RGBColor(0xC6, 0x28, 0x28)

_HEX_RE = re.compile(r"^#[0-9a-fA-F]{6}$")

_DEFAULT_RGB = {
    "green": GREEN,
    "yellow": YELLOW,
    "red": RED,
}


def hex_to_rgb(hex_color: str) -> RGBColor | None:
    raw = (hex_color or "").strip()
    if not _HEX_RE.fullmatch(raw):
        return None
    return RGBColor(int(raw[1:3], 16), int(raw[3:5], 16), int(raw[5:7], 16))


def resolve_highlight_palette(
    highlight: dict[str, Any] | None = None,
) -> dict[str, tuple[bool, RGBColor]]:
    """
    返回 green/yellow/red → (enabled, RGBColor)。
    enabled=False 时 Word 不着色该类；色值仍可用于网页。
    """
    out: dict[str, tuple[bool, RGBColor]] = {
        key: (True, rgb) for key, rgb in _DEFAULT_RGB.items()
    }
    if not isinstance(highlight, dict):
        return out
    for key in ("green", "yellow", "red"):
        item = highlight.get(key)
        enabled = True
        rgb = _DEFAULT_RGB[key]
        if isinstance(item, dict):
            raw_en = item.get("enabled", True)
            if isinstance(raw_en, bool):
                enabled = raw_en
            elif raw_en in (0, 1, "0", "1", "true", "false", "True", "False"):
                enabled = str(raw_en).lower() in ("1", "true")
            parsed = hex_to_rgb(str(item.get("hex") or ""))
            if parsed is not None:
                rgb = parsed
        out[key] = (enabled, rgb)
    return out


def _clear_runs(paragraph: Paragraph) -> None:
    for run in list(paragraph.runs):
        element = run._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def colorize_paragraph(
    paragraph: Paragraph,
    *,
    green: Iterable[str],
    yellow: Iterable[str],
    red: Iterable[str],
    highlight: dict[str, Any] | None = None,
) -> None:
    """按绿/黄/红三类为段落中的英语单词着色。"""
    text = paragraph.text
    if not text:
        return

    green_set = green if isinstance(green, (set, frozenset)) else set(green)
    yellow_set = yellow if isinstance(yellow, (set, frozenset)) else set(yellow)
    red_set = red if isinstance(red, (set, frozenset)) else set(red)
    palette = resolve_highlight_palette(highlight)

    _clear_runs(paragraph)

    last = 0
    for match in iter_english_tokens(text):
        if match.start() > last:
            paragraph.add_run(text[last : match.start()])

        token = match.group(0)
        key = token.lower()
        run = paragraph.add_run(token)
        if key in green_set:
            enabled, rgb = palette["green"]
            if enabled:
                run.font.color.rgb = rgb
        elif key in yellow_set:
            enabled, rgb = palette["yellow"]
            if enabled:
                run.font.color.rgb = rgb
        elif key in red_set:
            enabled, rgb = palette["red"]
            if enabled:
                run.font.color.rgb = rgb
        last = match.end()

    if last < len(text):
        paragraph.add_run(text[last:])


def _iter_all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
        if section.different_first_page_header_footer:
            yield from section.first_page_header.paragraphs
            yield from section.first_page_footer.paragraphs


def build_highlighted_docx(
    file_bytes: bytes,
    *,
    green: Iterable[str],
    yellow: Iterable[str],
    red: Iterable[str],
    highlight: dict[str, Any] | None = None,
) -> bytes:
    """生成三色标注后的 docx 字节内容。"""
    if not file_bytes:
        raise ValueError("上传文件内容为空")

    document = Document(io.BytesIO(file_bytes))
    green_set = green if isinstance(green, (set, frozenset)) else set(green)
    yellow_set = yellow if isinstance(yellow, (set, frozenset)) else set(yellow)
    red_set = red if isinstance(red, (set, frozenset)) else set(red)

    for paragraph in _iter_all_paragraphs(document):
        colorize_paragraph(
            paragraph,
            green=green_set,
            yellow=yellow_set,
            red=red_set,
            highlight=highlight,
        )

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
