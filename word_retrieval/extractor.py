"""从 Word 文档中提取英语单词。"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator

from docx import Document

# 字母词，或 1st/2nd 等序数缩写
ENGLISH_WORD_RE = re.compile(
    r"[A-Za-z]+(?:'[A-Za-z]+)?|\d+(?:st|nd|rd|th)",
    re.IGNORECASE,
)

# 有独立语义的单字母词（其余单字母如型号中的 V、o 不提取）
MEANINGFUL_SINGLE_LETTERS = frozenset({"a", "i"})


def is_meaningful_english_token(token: str) -> bool:
    """判断是否为应提取/着色的英语词元。"""
    if not token:
        return False
    # 序数缩写始终保留
    if token[0].isdigit():
        return True
    if len(token) == 1:
        return token.lower() in MEANINGFUL_SINGLE_LETTERS
    return True


def iter_english_tokens(text: str) -> Iterator[re.Match[str]]:
    """迭代文本中有意义的英语词元匹配。"""
    for match in ENGLISH_WORD_RE.finditer(text):
        if is_meaningful_english_token(match.group(0)):
            yield match


def extract_english_words(text: str) -> set[str]:
    """从文本中提取英语单词（小写去重）。"""
    return {match.group(0).lower() for match in iter_english_tokens(text)}


def read_docx_text(file_bytes: bytes) -> str:
    """读取 .docx 中的全部段落与表格文本。"""
    if not file_bytes:
        raise ValueError("上传文件内容为空")

    document = Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)
